# -*- coding: utf-8 -*-
"""
Rotina Smartsheet (via Excel exportado) - OLE VM - DEV IMOB
Lê o .xlsx exportado do Smartsheet, reconstrói a hierarquia UGB > Projeto > Setor
pelo recuo (indent) da coluna "Nome da tarefa" e separa atividades pendentes em:
atrasadas / semana / quinzena / mes. Gera atividades.json nesta pasta.

Uso: python parse_excel.py [caminho_do_xlsx]
Padrão: ..\..\06 Smartsheet\OLE VM - DEV IMOB.xlsx
"""
import json
import os
import re
import sys
from datetime import date, datetime, timedelta

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PASTA_XLSX = os.path.normpath(os.path.join(HERE, "..", "..", "06 Smartsheet"))


def newest_export():
    """Pega o 'OLE VM - DEV IMOB*.xlsx' mais recente da pasta 06 Smartsheet."""
    import glob
    cands = glob.glob(os.path.join(PASTA_XLSX, "*DEV IMOB*.xlsx"))
    if not cands:
        sys.exit(f"ERRO: nenhum 'OLE VM - DEV IMOB*.xlsx' em {PASTA_XLSX}")
    return max(cands, key=os.path.getmtime)


def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def parse_pct(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v) * 100 if v <= 1 else float(v)
    s = str(v).replace("%", "").replace(",", ".").strip()
    try:
        f = float(s)
        return f * 100 if f <= 1 else f
    except ValueError:
        return 0.0


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else newest_export()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    headers = {}
    for c in ws[1]:
        if c.value:
            headers[str(c.value).strip().lower()] = c.column

    def col(*needles):
        for h, idx in headers.items():
            if all(n in h for n in needles):
                return idx
        return None

    c_nome = col("nome")
    c_ini = col("iniciar") or col("início") or col("inicio")
    c_fim = col("terminar")
    c_pct = col("%")
    c_status = col("status")
    c_resp = col("resp")
    c_bloco = col("bloco")
    c_fim_lb = col("término lb") or col("termino lb")
    c_var = col("var")
    linha0 = 2  # dados comecam depois do cabecalho
    if not c_nome or not c_fim:
        # Export sem a linha de cabecalho (acontece quando o Smartsheet e
        # exportado por uma visao): assume a ordem fixa de colunas do
        # "OLE VM - DEV IMOB" e le desde a primeira linha.
        print("AVISO: export sem cabecalho — usando a ordem padrao de colunas.")
        c_bloco, c_resp, c_pct, c_nome = 1, 2, 4, 5
        c_ini, c_fim = 6, 7
        c_fim_lb, c_var, c_status = 10, 12, 13
        linha0 = 1

    # 1a passada: nome + indent de cada linha
    rows = []
    for r in range(linha0, ws.max_row + 1):
        cell = ws.cell(row=r, column=c_nome)
        raw = cell.value
        if raw is None or str(raw).strip() == "":
            continue
        name = str(raw)
        indent = cell.alignment.indent or 0
        # fallback: recuo por espacos no texto
        if indent == 0:
            m = re.match(r"^(\s+)", name)
            if m:
                indent = len(m.group(1)) // 2
        rows.append({
            "row": r,
            "name": name.strip(),
            "indent": int(indent),
            "ini": parse_date(ws.cell(row=r, column=c_ini).value) if c_ini else None,
            "fim": parse_date(ws.cell(row=r, column=c_fim).value),
            "pct": parse_pct(ws.cell(row=r, column=c_pct).value) if c_pct else 0.0,
            "status": ws.cell(row=r, column=c_status).value if c_status else None,
            "resp": ws.cell(row=r, column=c_resp).value if c_resp else None,
            "bloco": ws.cell(row=r, column=c_bloco).value if c_bloco else None,
            "fim_lb": parse_date(ws.cell(row=r, column=c_fim_lb).value) if c_fim_lb else None,
            "var": ws.cell(row=r, column=c_var).value if c_var else None,
        })

    # folha = linha cuja proxima linha nao tem indent maior
    for i, it in enumerate(rows):
        nxt = rows[i + 1]["indent"] if i + 1 < len(rows) else -1
        it["leaf"] = nxt <= it["indent"]

    hoje = date.today()
    fim_semana = hoje + timedelta(days=(6 - hoje.weekday()))
    fim_quinzena = hoje + timedelta(days=14)
    fim_mes = (date(hoje.year + (hoje.month == 12), (hoje.month % 12) + 1, 1)
               - timedelta(days=1))

    buckets = {"atrasadas": [], "semana": [], "quinzena": [], "mes": []}
    trilha = {}  # indent -> nome do ancestral atual

    for it in rows:
        trilha[it["indent"]] = it["name"]
        for k in [k for k in trilha if k > it["indent"]]:
            del trilha[k]
        if not it["leaf"] or it["fim"] is None or it["pct"] >= 100:
            continue

        ugb = trilha.get(0, "")
        projeto = trilha.get(1, "")
        setor = trilha.get(2, "(direto no projeto)") if it["indent"] > 2 else "(direto no projeto)"
        grupo = " > ".join(trilha[k] for k in sorted(trilha) if 2 < k < it["indent"])

        item = {
            "atividade": it["name"],
            "ugb": ugb,
            "projeto": projeto,
            "setor": setor if it["indent"] > 2 else (trilha.get(2, "") if it["indent"] == 3 else setor),
            "grupo": grupo,
            "responsavel": it["resp"],
            "inicio": it["ini"].isoformat() if it["ini"] else "",
            "termino": it["fim"].isoformat(),
            "pct": round(it["pct"]),
            "status": it["status"],
            "linha": it["row"],
        }
        # setor correto: ancestral de indent 2 quando a atividade esta abaixo dele
        item["setor"] = trilha.get(2, "(direto no projeto)") if it["indent"] > 2 else "(direto no projeto)"

        f = it["fim"]
        if f < hoje:
            buckets["atrasadas"].append(item)
        elif f <= fim_semana:
            buckets["semana"].append(item)
        elif f <= fim_quinzena:
            buckets["quinzena"].append(item)
        elif f <= fim_mes:
            buckets["mes"].append(item)

    for b in buckets.values():
        b.sort(key=lambda i: (i["ugb"], i["projeto"], i["setor"], i["termino"]))

    out = {
        "fonte": path,
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "hoje": hoje.isoformat(),
        "janelas": {
            "semana_ate": fim_semana.isoformat(),
            "quinzena_ate": fim_quinzena.isoformat(),
            "mes_ate": fim_mes.isoformat(),
        },
        "totais": {k: len(v) for k, v in buckets.items()},
        "atividades": buckets,
    }
    dest = os.path.join(HERE, "atividades.json")
    with open(dest, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=2)

    # ---- MARCOS (metas por setor) ----
    # A meta de cada marco e a data "Terminar" do PRIMEIRO export em que ele
    # apareceu (dia 0 = 14/09/2026, definido pela Carol). O baseline e
    # add-only: nunca sobrescreve metas ja registradas; as colunas de
    # baseline/variacao do proprio Smartsheet sao ignoradas de proposito.
    # Engenharia ficou de fora dos marcos de apresentacao (decisao da Carol em
    # 19/10/2026, quando a Engenharia foi reorganizada no Smartsheet).
    FORA_DOS_MARCOS = ("ENGENHARIA",)
    # UGB/projeto MODELO e um gabarito do macrofluxo (copia das metas do
    # Caucaia I) — nao e obra real, entao nao entra em marcos nem no baseline.
    # T+1 (Terra mais um) saiu do escopo em 19/10/2026: projeto congelado.
    # Para trazer de volta, basta remover o termo desta lista — as metas
    # continuam guardadas no marcos_baseline.json.
    UGB_FORA = ("MODELO", "MACROFLUXO", "T+1")
    marcos = []
    trilha_m = {}
    for it in rows:
        trilha_m[it["indent"]] = it["name"]
        for k in [k for k in trilha_m if k > it["indent"]]:
            del trilha_m[k]
        if str(it.get("bloco") or "").strip().lower() != "marco":
            continue
        ugb = trilha_m.get(0, "") if it["indent"] > 0 else ""
        projeto = trilha_m.get(1, "") if it["indent"] > 1 else ""
        setor = trilha_m.get(2, "") if it["indent"] > 2 else ""
        if any(p in str(setor).upper() for p in FORA_DOS_MARCOS):
            continue
        if any(p in (str(ugb) + " " + str(projeto)).upper() for p in UGB_FORA):
            continue
        mid = "|".join([ugb, projeto, it["name"]]).upper()
        vistos = sum(1 for m in marcos if m["id"] == mid or m["id"].startswith(mid + "#"))
        if vistos:
            mid = f"{mid}#{vistos + 1}"
        marcos.append({
            "id": mid,
            "marco": it["name"],
            "ugb": ugb, "projeto": projeto, "setor": setor,
            "inicio": it["ini"].isoformat() if it["ini"] else "",
            "termino_atual": it["fim"].isoformat() if it["fim"] else "",
            "pct": round(it["pct"]), "status": it["status"],
            "responsavel": it["resp"], "linha": it["row"],
        })

    base_path = os.path.join(HERE, "marcos_baseline.json")
    baseline = {}
    if os.path.exists(base_path):
        with open(base_path, encoding="utf-8") as fh:
            baseline = json.load(fh)
    novos = 0
    for m in marcos:
        if m["id"] not in baseline and m["termino_atual"]:
            baseline[m["id"]] = {"meta": m["termino_atual"],
                                 "registrado_em": hoje.isoformat()}
            novos += 1
    with open(base_path, "w", encoding="utf-8") as fh:
        json.dump(baseline, fh, ensure_ascii=False, indent=2)

    for m in marcos:
        b = baseline.get(m["id"])
        m["meta"] = b["meta"] if b else m["termino_atual"]
        m["meta_registrada_em"] = b["registrado_em"] if b else hoje.isoformat()
        if m["meta"] and m["termino_atual"]:
            d1 = datetime.strptime(m["meta"], "%Y-%m-%d").date()
            d2 = datetime.strptime(m["termino_atual"], "%Y-%m-%d").date()
            m["desvio_dias"] = (d2 - d1).days
        else:
            m["desvio_dias"] = None

    dest_m = os.path.join(HERE, "marcos.json")
    with open(dest_m, "w", encoding="utf-8") as fh:
        json.dump({"gerado_em": datetime.now().isoformat(timespec="seconds"),
                   "hoje": hoje.isoformat(), "marcos": marcos},
                  fh, ensure_ascii=False, indent=2)
    print(f"Marcos: {len(marcos)} (novos no baseline: {novos})")

    print("Fonte:", path)
    print("Linhas lidas:", len(rows))
    for k, v in out["totais"].items():
        print(f"  {k}: {v}")
    print("Gerado:", dest)


if __name__ == "__main__":
    main()
