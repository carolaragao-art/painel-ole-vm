# -*- coding: utf-8 -*-
"""Lê os projetos de engenharia do Condomínio Morada dos Sabiás direto do
export do Smartsheet (OLE VM - DEV IMOB*.xlsx) e devolve, para cada projeto,
a etapa atual e a data planejada (Terminar) daquela etapa.

Fonte da verdade: Smartsheet (cronograma). Os dados de acompanhamento
(empresa, projetista, pagamento) vêm do de-para em depara_projetos.json.
"""
import hashlib
import json
import os
import re
import unicodedata

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
RAIZ = "CONDOMÍNIO MORADA DOS SABIÁS"
GALHO = "ENGENHARIA DE PRODUTO"

# Ordem canônica do fluxo. As 4 primeiras vêm do Smartsheet; Pagamento,
# Órgão e Aprovado só existem no painel (etapas 5 a 7).
ETAPAS = ["Contratação", "Projeto", "Análise crítica", "Finalização",
          "Pagamento", "Órgão", "Aprovado"]
# palavra-chave -> índice da etapa (0-based)
PADROES = [
    (r"contrata|alinhamento com o projetista|or[çc]amento", 0),
    (r"elabora|revis[ãa]o do projeto|atualiza[çc][ãa]o do projeto|execu[çc][ãa]o", 1),
    (r"an[áa]lise cr[íi]tica", 2),
    (r"finaliza", 3),
]


def ident(caminho):
    """ID estável do projeto: md5 do caminho normalizado. Precisa ser igual
    entre execuções, senão as edições feitas no painel perdem o vínculo."""
    base = chave(" > ".join(caminho))
    return "SS" + hashlib.md5(base.encode("utf-8")).hexdigest()[:10]


def chave(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def etapa_do_nome(nome):
    k = chave(nome)
    for padrao, idx in PADROES:
        if re.search(padrao, k):
            return idx
    return None


def iso(v):
    return v.date().isoformat() if hasattr(v, "date") else (str(v)[:10] if v else "")


def pct(v):
    try:
        return round(float(v) * 100)
    except (TypeError, ValueError):
        return None


def ler_linhas(xlsx):
    """Devolve as linhas do galho do condomínio, com caminho e indentação."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    primeira = ws.cell(row=1, column=5).value
    r0 = 2 if isinstance(primeira, str) and "nome" in primeira.lower() else 1
    pilha, fora = {}, []
    for r in ws.iter_rows(min_row=r0):
        nome = r[4].value
        if not isinstance(nome, str) or not nome.strip():
            continue
        ind = int(r[4].alignment.indent or 0)
        pilha[ind] = nome.strip()
        for k in [k for k in list(pilha) if k > ind]:
            del pilha[k]
        if chave(RAIZ) not in chave(pilha.get(1, "")) or chave(GALHO) not in chave(pilha.get(2, "")):
            continue
        fora.append({
            "ind": ind, "nome": nome.strip(),
            "caminho": [pilha[k] for k in sorted(pilha) if k <= ind],
            "pct": pct(r[3].value), "fim": iso(r[6].value),
        })
    wb.close()
    return fora


def extrair(xlsx):
    """Projeto = nó cujos filhos diretos são etapas do fluxo."""
    linhas = ler_linhas(xlsx)
    projetos, avisos = [], []
    # % que o próprio Smartsheet calcula para o galho e para cada grupo
    rollup = {"geral": None, "grupos": {}}
    for l in linhas:
        if l["ind"] == 2:
            rollup["geral"] = l["pct"]
        elif l["ind"] == 3:
            rollup["grupos"][l["nome"]] = l["pct"]
    for i, l in enumerate(linhas):
        filhos = []
        for prox in linhas[i + 1:]:
            if prox["ind"] <= l["ind"]:
                break
            if prox["ind"] == l["ind"] + 1:
                filhos.append(prox)
        etapas = [(etapa_do_nome(f["nome"]), f) for f in filhos]
        reconhecidas = [(e, f) for e, f in etapas if e is not None]
        if len(reconhecidas) < 2:
            # Item simples: filho direto de um grupo, sem etapas detalhadas
            # (ex.: TOPOGRAFIA, SONDAGEM POÇO PROFUNDO). Entra pelo próprio %.
            if l["ind"] == 4 and not filhos and l["pct"] is not None:
                p = l["pct"]
                projetos.append({
                    "id": ident(l["caminho"]),
                    "caminho": l["caminho"][3:], "grupo": l["caminho"][3],
                    "nome": l["nome"], "pct": p,
                    "etapa": 4 if p >= 100 else (0 if p == 0 else 1),
                    "etapa_nome": "Pronto para pagamento" if p >= 100 else ("Contratação" if p == 0 else "Projeto"),
                    "passo": "sem etapas detalhadas no Smartsheet",
                    "prazo_plano": l["fim"], "fim_projeto": l["fim"],
                    "sub": [], "atipico": True,
                })
            continue
        if len(reconhecidas) < len(filhos):
            nao = [f["nome"] for e, f in etapas if e is None]
            avisos.append({"projeto": l["nome"], "sub_etapas_fora_do_padrao": nao})
        reconhecidas.sort(key=lambda x: x[0])
        pendentes = [(e, f) for e, f in reconhecidas if (f["pct"] or 0) < 100]
        if pendentes:
            etapa, passo = pendentes[0]
            concluido = False
        else:
            etapa, passo = reconhecidas[-1][0], reconhecidas[-1][1]
            concluido = True
        grupo = l["caminho"][3] if len(l["caminho"]) > 3 else ""
        projetos.append({
            "id": ident(l["caminho"]),
            "caminho": l["caminho"][3:],          # a partir do grupo
            "grupo": grupo,
            "nome": l["nome"],
            "pct": l["pct"],
            "etapa": 4 if concluido else etapa,   # tudo pronto → aguarda pagamento
            "etapa_nome": "Pronto para pagamento" if concluido else ETAPAS[etapa],
            "passo": passo["nome"],
            "prazo_plano": passo["fim"],
            "fim_projeto": l["fim"],
            "sub": [{"nome": f["nome"], "pct": f["pct"], "fim": f["fim"],
                     "etapa": e} for e, f in reconhecidas],
            "atipico": len(reconhecidas) < 3,
        })
    return projetos, avisos, rollup


if __name__ == "__main__":
    import sys, glob
    pasta = os.path.normpath(os.path.join(HERE, "..", "..", "06 Smartsheet"))
    xlsx = sys.argv[1] if len(sys.argv) > 1 else max(
        glob.glob(os.path.join(pasta, "*DEV IMOB*.xlsx")), key=os.path.getmtime)
    projs, avisos, _ = extrair(xlsx)
    print("Fonte:", os.path.basename(xlsx))
    print("Projetos encontrados:", len(projs))
    for p in projs:
        print("  [%-22s] %-44s etapa %d · %-22s prazo %s%s" % (
            p["grupo"][:22], " > ".join(p["caminho"][1:])[:44] or p["nome"][:44],
            p["etapa"] + 1, p["etapa_nome"], p["prazo_plano"],
            "  ATÍPICO" if p["atipico"] else ""))
    if avisos:
        print("\nSub-etapas fora do padrão:")
        for a in avisos:
            print("  -", a["projeto"], "->", a["sub_etapas_fora_do_padrao"])
    print(json.dumps({"n": len(projs)}, ensure_ascii=False))
